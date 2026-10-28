"""Generate operator-ready PowerPoint reports from reviewed results."""

from __future__ import annotations

import getpass
import io
import os
import re
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .inference import InferenceResult
from .rendering import RenderOptions, render_result


@dataclass(frozen=True)
class ReportImage:
    image: Image.Image


@dataclass(frozen=True)
class ReportSlide:
    visual_id: str
    images: tuple[ReportImage, ...]


@dataclass(frozen=True)
class PowerPointReport:
    path: Path
    author: str
    image_count: int
    slide_count: int


@dataclass(frozen=True)
class _Box:
    left: float
    top: float
    width: float
    height: float

    @property
    def right(self) -> float:
        return self.left + self.width

    @property
    def bottom(self) -> float:
        return self.top + self.height


@dataclass(frozen=True)
class _ImagePlacement:
    image_box: _Box
    description_box: _Box


@dataclass(frozen=True)
class _SlideLayout:
    information_position: str
    header_box: _Box
    information_box: _Box
    placements: tuple[_ImagePlacement, ...]
    image_area: float


_HEADER = _Box(0.12, 0.12, 13.09, 0.7)
_SIDE_INFORMATION = _Box(0.12, 0.95, 2.5, 6.43)
_SIDE_IMAGES = _Box(2.75, 0.95, 10.46, 6.43)
_BOTTOM_IMAGES = _Box(0.12, 0.95, 13.09, 4.89)
_BOTTOM_INFORMATION = _Box(0.12, 5.97, 13.09, 1.41)
_CELL_GAP = 0.12
_DESCRIPTION_GAP = 0.06
_DESCRIPTION_HEIGHT = 0.3
_IMAGE_DESCRIPTION = "[Enter image description]"


def visual_id_from_path(path: Path) -> str:
    explicit = re.search(r"Z\d+", path.stem, flags=re.IGNORECASE)
    if explicit:
        return explicit.group(0).upper()
    trailing_digits = re.search(r"(\d{5})$", path.stem)
    if trailing_digits:
        return f"Z{trailing_digits.group(1)}"
    return path.stem


def powerpoint_author() -> str:
    """Return the user name configured for Microsoft Office."""
    for key_path in (
        r"Software\Microsoft\Office\Common\UserInfo",
        r"Software\Microsoft\Office\16.0\Common\UserInfo",
    ):
        value = _read_registry_value(key_path, "UserName")
        if value:
            return value
    return getpass.getuser()


def _read_registry_value(key_path: str, value_name: str) -> str | None:
    if os.name != "nt":
        return None
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
            value, _value_type = winreg.QueryValueEx(key, value_name)
    except OSError:
        return None
    text = str(value).strip()
    return text or None


def combine_report_images(images: list[tuple[str, Image.Image]]) -> Image.Image:
    """Combine model outputs for one frame without changing their aspect ratios."""
    if not images:
        raise ValueError("At least one rendered model image is required.")
    if len(images) == 1:
        return images[0][1].convert("RGB")

    header_height = 34
    gap = 12
    normalized = [image.convert("RGB") for _label, image in images]
    height = max(image.height for image in normalized)
    widths = [round(image.width * height / image.height) for image in normalized]
    combined = Image.new("RGB", (sum(widths) + gap * (len(images) - 1), height + header_height), "white")
    draw = ImageDraw.Draw(combined)
    font = ImageFont.truetype("calibri.ttf", 18) if Path(r"C:\Windows\Fonts\calibri.ttf").exists() else ImageFont.load_default()
    left = 0
    for (label, _source), image, width in zip(images, normalized, widths):
        resized = image.resize((width, height), Image.Resampling.LANCZOS)
        combined.paste(resized, (left, header_height))
        draw.text((left + width / 2, 7), label, fill="black", font=font, anchor="ma")
        left += width + gap
    return combined


def render_report_image(
    source: Image.Image,
    options: RenderOptions,
    models: list[tuple[str, InferenceResult, set[int]]],
) -> Image.Image:
    rendered = [
        (label, render_result(source, result, options, hidden_annotations=hidden))
        for label, result, hidden in models
    ]
    return combine_report_images(rendered)


def generate_powerpoint_report(
    destination: Path,
    slides: list[ReportSlide],
    author: str | None = None,
) -> PowerPointReport:
    if not slides:
        raise ValueError("At least one VID is required for the report.")
    if any(not 1 <= len(slide.images) <= 3 for slide in slides):
        raise ValueError("Select between one and three result frames per VID.")
    destination = destination.with_suffix(".pptx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    report_author = (author or powerpoint_author()).strip() or getpass.getuser()

    presentation = Presentation()
    presentation.slide_width = 12_192_000
    presentation.slide_height = 6_858_000
    presentation.core_properties.author = report_author
    presentation.core_properties.last_modified_by = report_author
    presentation.core_properties.title = "EchoSight Inspection Report"
    presentation.core_properties.subject = ", ".join(slide.visual_id for slide in slides)
    presentation.core_properties.comments = "Generated by EchoSight 2.0 from operator-reviewed active annotations."

    for report_slide in slides:
        _add_report_slide(presentation, report_slide)

    presentation.save(destination)
    return PowerPointReport(destination, report_author, sum(len(slide.images) for slide in slides), len(slides))


def _add_report_slide(presentation: Presentation, report: ReportSlide) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = _rgb(255, 255, 255)

    layout = _choose_slide_layout(report.images)
    _add_information_panel(slide, report.visual_id, layout.information_position)
    for item, placement in zip(report.images, layout.placements):
        image = item.image.convert("RGB")
        stream = io.BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
        image_box = placement.image_box
        slide.shapes.add_picture(
            stream,
            Inches(image_box.left),
            Inches(image_box.top),
            width=Inches(image_box.width),
            height=Inches(image_box.height),
        )
        description_box = placement.description_box
        _add_highlighted_text(
            slide,
            _IMAGE_DESCRIPTION,
            description_box.left,
            description_box.top,
            description_box.width,
            description_box.height,
            9,
            alignment=PP_ALIGN.CENTER,
        )


def _choose_slide_layout(images: tuple[ReportImage, ...]) -> _SlideLayout:
    candidates = (
        _layout_for_region(images, "side", _SIDE_INFORMATION, _SIDE_IMAGES),
        _layout_for_region(images, "bottom", _BOTTOM_INFORMATION, _BOTTOM_IMAGES),
    )
    return max(candidates, key=lambda candidate: candidate.image_area)


def _layout_for_region(
    images: tuple[ReportImage, ...],
    information_position: str,
    information_box: _Box,
    image_region: _Box,
) -> _SlideLayout:
    row_options = {
        1: ((1,),),
        2: ((2,), (1, 1)),
        3: ((3,), (1, 2), (2, 1), (1, 1, 1)),
    }[len(images)]
    layouts = []
    for row_counts in row_options:
        row_height = (image_region.height - _CELL_GAP * (len(row_counts) - 1)) / len(row_counts)
        placements = []
        image_index = 0
        for row_index, column_count in enumerate(row_counts):
            column_width = (image_region.width - _CELL_GAP * (column_count - 1)) / column_count
            cell_top = image_region.top + row_index * (row_height + _CELL_GAP)
            for column_index in range(column_count):
                cell_left = image_region.left + column_index * (column_width + _CELL_GAP)
                available = _Box(
                    cell_left,
                    cell_top,
                    column_width,
                    row_height - _DESCRIPTION_GAP - _DESCRIPTION_HEIGHT,
                )
                image_box = _fit_image_box(images[image_index].image, available)
                description_width = min(column_width, _compact_text_width(_IMAGE_DESCRIPTION, 9))
                description_box = _Box(
                    cell_left + (column_width - description_width) / 2,
                    cell_top + row_height - _DESCRIPTION_HEIGHT,
                    description_width,
                    _DESCRIPTION_HEIGHT,
                )
                placements.append(_ImagePlacement(image_box, description_box))
                image_index += 1
        image_area = sum(placement.image_box.width * placement.image_box.height for placement in placements)
        layouts.append(_SlideLayout(information_position, _HEADER, information_box, tuple(placements), image_area))
    return max(layouts, key=lambda layout: layout.image_area)


def _fit_image_box(image: Image.Image, available: _Box) -> _Box:
    scale = min(available.width / image.width, available.height / image.height)
    width = image.width * scale
    height = image.height * scale
    return _Box(
        available.left + (available.width - width) / 2,
        available.top + (available.height - height) / 2,
        width,
        height,
    )


def _boxes_overlap(first: _Box, second: _Box) -> bool:
    return first.left < second.right and second.left < first.right and first.top < second.bottom and second.top < first.bottom


def _add_information_panel(slide: object, visual_id: str, position: str) -> None:
    sighting = "Sighting: [Enter sighting]"
    sighting_width = _compact_text_width(sighting, 10)
    _add_highlighted_text(slide, sighting, 0.2, 0.12, sighting_width, 0.36, 10)
    visual_id_text = f"Visual ID: {visual_id}"
    visual_id_left = 0.2 + sighting_width + 0.2
    visual_id_width = _compact_text_width(visual_id_text, 10, maximum=13.13 - visual_id_left)
    _add_highlighted_text(slide, visual_id_text, visual_id_left, 0.12, visual_id_width, 0.36, 10)

    comments = "Comments: [Enter comments]"
    if position == "side":
        _add_highlighted_text(slide, comments, 0.2, 1.02, _compact_text_width(comments, 9, maximum=2.3), 0.36, 9)
        _add_verification_notice(slide, 0.2, 6.78, 2.3, 0.48)
        return

    _add_highlighted_text(slide, comments, 0.2, 6.02, _compact_text_width(comments, 9), 0.36, 9)
    _add_verification_notice(slide, 9.55, 6.15, 3.58, 0.72)


def _compact_text_width(text: str, font_size: int, maximum: float = 6.0) -> float:
    estimated_width = len(text) * font_size / 72 * 0.5 + 0.18
    return min(maximum, max(1.0, estimated_width))


def _add_verification_notice(slide: object, left: float, top: float, width: float, height: float) -> None:
    _add_highlighted_text(
        slide,
        "* This is AI analyzed data, Engineer to confirm results; delete this text field after results are validated.",
        left,
        top,
        width,
        height,
        5,
    )


def _add_highlighted_text(
    slide: object,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    font_size: int,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(255, 242, 0)
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = alignment
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    run.font.color.rgb = _rgb(0, 0, 0)


def _add_text(
    slide: object,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    font_size: int,
    bold: bool = False,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = alignment
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(0, 0, 0)


def _rgb(red: int, green: int, blue: int) -> object:
    from pptx.dml.color import RGBColor

    return RGBColor(red, green, blue)