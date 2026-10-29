from pathlib import Path
from unittest.mock import patch

import pytest
from echosight2.inference import InferenceResult, TaskType
from echosight2.rendering import RenderOptions
from echosight2.reporting import (
    ReportImage,
    ReportSlide,
    _boxes_overlap,
    _choose_slide_layout,
    combine_report_images,
    generate_powerpoint_report,
    powerpoint_author,
    render_report_image,
    visual_id_from_path,
)
from PIL import Image
from pptx import Presentation


def test_powerpoint_author_prefers_office_profile() -> None:
    with patch("echosight2.reporting._read_registry_value", side_effect=("Office Operator", None)):
        assert powerpoint_author() == "Office Operator"


def test_powerpoint_author_falls_back_to_windows_user() -> None:
    with patch("echosight2.reporting._read_registry_value", return_value=None), patch(
        "echosight2.reporting.getpass.getuser", return_value="Windows Operator"
    ):
        assert powerpoint_author() == "Windows Operator"


def test_combines_dual_model_images_with_headers() -> None:
    combined = combine_report_images(
        [
            ("Detection", Image.new("RGB", (80, 120), "black")),
            ("Anomaly", Image.new("RGB", (60, 120), "gray")),
        ]
    )

    assert combined.width == 152
    assert combined.height == 154


def test_report_rendering_excludes_hidden_annotations_without_selection_highlight() -> None:
    source = Image.new("RGB", (80, 60), "black")
    result = InferenceResult(Path("inspection.tiff"), TaskType.DETECTION, (80, 60), (32, 32), 4.0)
    rendered = Image.new("RGB", source.size, "white")

    with patch("echosight2.reporting.render_result", return_value=rendered) as render:
        output = render_report_image(source, RenderOptions(), [("Detection", result, {1, 3})])

    assert output.tobytes() == rendered.tobytes()
    render.assert_called_once_with(source, result, RenderOptions(), hidden_annotations={1, 3})


def test_visual_id_comes_from_explicit_or_trailing_filename_value() -> None:
    assert visual_id_from_path(Path("scan_Z00123_gate.tiff")) == "Z00123"
    assert visual_id_from_path(Path("D6N07R7400095.tiff")) == "Z00095"
    assert visual_id_from_path(Path("unit_alpha.tiff")) == "unit_alpha"


@pytest.mark.parametrize(
    ("size", "expected_position"),
    [((1600, 400), "bottom"), ((400, 1600), "side")],
)
def test_report_layout_places_information_where_images_are_largest(
    size: tuple[int, int], expected_position: str
) -> None:
    layout = _choose_slide_layout((ReportImage(Image.new("RGB", size)),))

    assert layout.information_position == expected_position


@pytest.mark.parametrize("image_count", [1, 2, 3])
def test_report_layout_boxes_do_not_overlap(image_count: int) -> None:
    images = tuple(ReportImage(Image.new("RGB", (1200, 700))) for _index in range(image_count))
    layout = _choose_slide_layout(images)
    boxes = [layout.header_box, layout.information_box]
    for placement in layout.placements:
        boxes.extend((placement.image_box, placement.description_box))

    assert len(layout.placements) == image_count
    assert all(box.left >= 0 and box.top >= 0 and box.right <= 13.3334 and box.bottom <= 7.5 for box in boxes)
    assert all(box.bottom <= 6.851 for box in boxes)
    assert all(not _boxes_overlap(first, second) for index, first in enumerate(boxes) for second in boxes[index + 1 :])


def test_report_header_is_always_at_the_top() -> None:
    landscape = _choose_slide_layout((ReportImage(Image.new("RGB", (1600, 400))),))
    portrait = _choose_slide_layout((ReportImage(Image.new("RGB", (400, 1600))),))

    assert landscape.header_box == portrait.header_box
    assert landscape.header_box.top == pytest.approx(0.12)
    assert all(placement.image_box.top >= landscape.header_box.bottom for placement in landscape.placements)
    assert all(placement.image_box.top >= portrait.header_box.bottom for placement in portrait.placements)


def test_generates_one_formatted_slide_per_vid(tmp_path: Path) -> None:
    destination = tmp_path / "inspection_report"
    items = [ReportImage(Image.new("RGB", (100, 220), color)) for color in ("black", "gray", "white")]
    slides = [
        ReportSlide("Z00095", tuple(items)),
        ReportSlide("Z00102", (items[0],)),
    ]

    report = generate_powerpoint_report(destination, slides, author="Office Operator")

    assert report.path == destination.with_suffix(".pptx")
    assert report.author == "Office Operator"
    assert report.image_count == 4
    assert report.slide_count == 2
    presentation = Presentation(report.path)
    assert presentation.slide_width == 12_192_000
    assert presentation.slide_height == 6_858_000
    assert presentation.core_properties.author == "Office Operator"
    assert presentation.core_properties.title == "EchoSight Inspection Report"
    assert len(presentation.slides) == 2
    assert presentation.slides[0].slide_layout.name == "Title and Content"
    assert not presentation.slides[0].placeholders
    master_text = "\n".join(shape.text for shape in presentation.slide_masters[0].shapes if hasattr(shape, "text_frame"))
    assert "PQE SID JFDL FIFA" not in master_text
    assert "Intel Confidential" in master_text
    assert sum(1 for shape in presentation.slides[0].shapes if shape.shape_type == 13) == 3
    first_slide_fields = [shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text_frame")]
    first_slide_text = "\n".join(first_slide_fields)
    assert "Sighting: [Enter sighting]" in first_slide_fields
    assert "Visual ID: Z00095" in first_slide_fields
    assert "Comments: [Enter comments]" in first_slide_fields
    assert "PQE SID JFDL FIFA" in first_slide_fields
    assert "Sighting:" not in first_slide_fields
    assert "Visual ID:" not in first_slide_fields
    assert "Comments:" not in first_slide_fields
    assert first_slide_text.count("[Enter image description]") == 3
    assert "Gate" not in first_slide_text
    assert "delete this text field" in first_slide_text
    editable_fields = {
        shape.text: shape
        for shape in presentation.slides[0].shapes
        if hasattr(shape, "text_frame") and shape.text in {"Sighting: [Enter sighting]", "Visual ID: Z00095", "Comments: [Enter comments]"}
    }
    assert editable_fields["Sighting: [Enter sighting]"].width / 914_400 < 2.2
    assert editable_fields["Visual ID: Z00095"].width / 914_400 < 1.6
    assert editable_fields["Comments: [Enter comments]"].width / 914_400 < 2.0
    description_widths = [
        shape.width / 914_400
        for shape in presentation.slides[0].shapes
        if hasattr(shape, "text_frame") and shape.text == "[Enter image description]"
    ]
    assert all(width < 1.8 for width in description_widths)
    for slide in presentation.slides:
        department = next(shape for shape in slide.shapes if getattr(shape, "text", "") == "PQE SID JFDL FIFA")
        assert not department.is_placeholder
        assert department.left / 914_400 == pytest.approx(0.7, abs=0.001)
        assert department.top / 914_400 == pytest.approx(7.125)


def test_report_falls_back_when_template_is_unavailable(tmp_path: Path) -> None:
    destination = tmp_path / "fallback_report.pptx"
    slides = [ReportSlide("Z00095", (ReportImage(Image.new("RGB", (100, 100), "black")),))]

    report = generate_powerpoint_report(destination, slides, author="Office Operator", template_path=tmp_path / "missing.pptx")

    presentation = Presentation(report.path)
    master_text = "\n".join(shape.text for shape in presentation.slide_masters[0].shapes if hasattr(shape, "text_frame"))
    assert "PQE SID JFDL FIFA" not in master_text
    assert any(getattr(shape, "text", "") == "PQE SID JFDL FIFA" for shape in presentation.slides[0].shapes)
    assert len(presentation.slides) == 1


def test_report_requires_slides_with_one_to_three_images(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="At least one VID"):
        generate_powerpoint_report(tmp_path / "empty.pptx", [])

    images = [ReportImage(Image.new("RGB", (10, 10))) for _index in range(4)]
    slide = ReportSlide("Z00095", tuple(images))
    with pytest.raises(ValueError, match="per VID"):
        generate_powerpoint_report(tmp_path / "too_many.pptx", [slide])